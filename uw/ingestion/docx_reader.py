"""Extract text from DOCX files."""
from pathlib import Path
import logging

logger = logging.getLogger(__name__)


def read_docx(path: Path) -> str:
    try:
        from docx import Document
        doc = Document(str(path))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                row_text = "\t".join(cell.text.strip() for cell in row.cells)
                if row_text.strip():
                    paragraphs.append(row_text)
        return "\n".join(paragraphs)
    except Exception as e:
        logger.error(f"DOCX read error {path.name}: {e}")
        return ""
